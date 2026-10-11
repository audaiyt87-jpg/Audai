# -*- coding: utf-8 -*-
"""
حاسبة الراتب — تطبيق ويندوز ثنائي اللغة (CustomTkinter).
Salary Calculator - modern bilingual (Arabic / English) Windows app.
الامتداد .pyw يفتح البرنامج بدون نافذة سوداء (console).
"""
import os
import sys
import html
import tempfile
import traceback
import webbrowser
import datetime
from tkinter import filedialog, messagebox

try:
    import customtkinter as ctk
except Exception:
    import tkinter as _tk
    _r = _tk.Tk()
    _r.withdraw()
    messagebox.showerror(
        "Missing library \u2014 \u0645\u0643\u062a\u0628\u0629 \u0646\u0627\u0642\u0635\u0629",
        "CustomTkinter is not installed.\n"
        "\u0627\u0644\u0645\u0643\u062a\u0628\u0629 \u063a\u064a\u0631 \u0645\u062b\u0628\u062a\u0629.\n\n"
        "\u0627\u0641\u062a\u062d Command Prompt \u0648\u0646\u0641\u0651\u0630:\n\n"
        "    pip install customtkinter openpyxl\n\n"
        "\u062b\u0645 \u0634\u063a\u0651\u0644 \u0627\u0644\u0628\u0631\u0646\u0627\u0645\u062c \u0645\u0631\u0629 \u0623\u062e\u0631\u0649.")
    sys.exit(1)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    _HAS_XLSX = True
except Exception:
    _HAS_XLSX = False

DAYS = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
FRI = [4, 4, 5, 4, 4, 5, 4, 4, 5, 4, 4, 5]
AR = "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669"
A2W = {AR[i]: str(i) for i in range(10)}

# key: (arabic, english)
T = {
    "app": ("\u062d\u0627\u0633\u0628\u0629 \u0627\u0644\u0631\u0627\u062a\u0628", "Salary Calculator"),
    "tin": ("\u0627\u0644\u0645\u062f\u062e\u0644\u0627\u062a", "Inputs"),
    "tout": ("\u0627\u0644\u0646\u062a\u064a\u062c\u0629", "Result"),
    "s_period": ("\u0627\u0644\u0641\u062a\u0631\u0629 \u0648\u0627\u0644\u0631\u0627\u062a\u0628 \u0627\u0644\u0623\u0633\u0627\u0633\u064a", "Period & Basic Salary"),
    "month": ("\u0627\u0644\u0634\u0647\u0631", "Month"),
    "basic": ("\u0627\u0644\u0631\u0627\u062a\u0628 \u0627\u0644\u0623\u0633\u0627\u0633\u064a", "Basic Salary"),
    "s_allow": ("\u0627\u0644\u0628\u062f\u0644\u0627\u062a", "Allowances"),
    "road_q": ("\u0647\u0644 \u064a\u0648\u062c\u062f \u0628\u062f\u0644 \u0637\u0631\u064a\u0642\u061f", "Road allowance?"),
    "road_v": ("\u0645\u0642\u062f\u0627\u0631 \u0628\u062f\u0644 \u0627\u0644\u0637\u0631\u064a\u0642 (0.07 / 0.15 / \u0633\u0627\u0639\u0627\u062a)",
               "Road allowance (0.07 / 0.15 / hours)"),
    "house_q": ("\u0647\u0644 \u064a\u0648\u062c\u062f \u0628\u062f\u0644 \u0633\u0643\u0646\u061f", "Housing allowance?"),
    "contract": ("\u0646\u0648\u0639 \u0627\u0644\u0639\u0642\u062f (\u0627\u0644\u0633\u0643\u0646)", "Contract type (housing)"),
    "c_fam": ("\u0639\u0627\u0626\u0644\u064a (3)", "Family (3)"),
    "c_sin": ("\u0623\u0639\u0632\u0628 (2)", "Single (2)"),
    "ot": ("\u0639\u062f\u062f \u0633\u0627\u0639\u0627\u062a \u0627\u0644\u0639\u0645\u0644 \u0627\u0644\u0625\u0636\u0627\u0641\u064a", "Overtime hours"),
    "yes": ("\u0646\u0639\u0645", "Yes"),
    "no": ("\u0644\u0627", "No"),
    "s_ded": ("\u0627\u0644\u062e\u0635\u0645\u064a\u0627\u062a", "Deductions"),
    "adv": ("\u0645\u0642\u062f\u0627\u0631 \u0627\u0644\u0633\u0644\u0641\u0629 (\u0625\u0646 \u0648\u062c\u062f)", "Advance amount (if any)"),
    "med": ("\u0645\u0642\u062f\u0627\u0631 \u0627\u0644\u062a\u0623\u0645\u064a\u0646 \u0627\u0644\u0637\u0628\u064a (\u0625\u0646 \u0648\u062c\u062f)", "Medical insurance (if any)"),
    "s_leave": ("\u062e\u0635\u0645\u064a\u0627\u062a \u0627\u0644\u0625\u062c\u0627\u0632\u0627\u062a \u0648\u0627\u0644\u063a\u064a\u0627\u0628", "Leave & Absence"),
    "lp": ("\u0623\u064a\u0627\u0645 \u0627\u0644\u0625\u062c\u0627\u0632\u0629 \u0645\u0646 \u0627\u0644\u0631\u0635\u064a\u062f", "Paid leave days"),
    "lu": ("\u0623\u064a\u0627\u0645 \u0627\u0644\u0625\u062c\u0627\u0632\u0629 \u0628\u062f\u0648\u0646 \u0631\u0635\u064a\u062f", "Unpaid leave days"),
    "absn": ("\u0623\u064a\u0627\u0645 \u0627\u0644\u063a\u064a\u0627\u0628 \u0628\u062f\u0648\u0646 \u0639\u0630\u0631", "Unexcused absence days"),
    "fing": ("\u0627\u0644\u0628\u0635\u0645\u0627\u062a \u0641\u0648\u0642 \u0627\u0644\u062d\u062f \u0627\u0644\u0645\u0633\u0645\u0648\u062d", "Fingerprints over limit"),
    "other": ("\u062e\u0635\u0645\u064a\u0627\u062a \u0623\u062e\u0631\u0649", "Other deductions"),
    "res_t": ("\u062a\u0641\u0627\u0635\u064a\u0644 \u0627\u0644\u0631\u0627\u062a\u0628", "Salary Breakdown"),
    "r_days": ("\u0639\u062f\u062f \u0623\u064a\u0627\u0645 \u0627\u0644\u0634\u0647\u0631", "Days in month"),
    "r_work": ("\u0623\u064a\u0627\u0645 \u0627\u0644\u062f\u0648\u0627\u0645 \u0627\u0644\u0641\u0639\u0644\u064a", "Actual working days"),
    "r_day": ("\u0642\u064a\u0645\u0629 \u064a\u0648\u0645 \u0627\u0644\u0639\u0645\u0644", "Work day value"),
    "r_hour": ("\u0642\u064a\u0645\u0629 \u0633\u0627\u0639\u0629 \u0627\u0644\u0639\u0645\u0644", "Work hour value"),
    "r_basic": ("\u0627\u0644\u0623\u0633\u0627\u0633\u064a", "Basic"),
    "r_fixed": ("\u0627\u0644\u062b\u0627\u0628\u062a", "Fixed"),
    "r_road": ("\u0625\u0636\u0627\u0641\u064a \u0628\u062f\u0644 \u0627\u0644\u0637\u0631\u064a\u0642", "Road allowance"),
    "r_house": ("\u0628\u062f\u0644 \u0627\u0644\u0633\u0643\u0646", "Housing allowance"),
    "r_ot": ("\u0627\u0644\u0639\u0645\u0644 \u0627\u0644\u0625\u0636\u0627\u0641\u064a", "Overtime"),
    "r_adv": ("\u062e\u0635\u0645 \u0627\u0644\u0633\u0644\u0641\u0629", "Advance deduction"),
    "r_med": ("\u062e\u0635\u0645 \u0627\u0644\u062a\u0623\u0645\u064a\u0646 \u0627\u0644\u0637\u0628\u064a", "Medical deduction"),
    "r_lp": ("\u062e\u0635\u0645 \u0627\u0644\u0625\u062c\u0627\u0632\u0629 \u0645\u0646 \u0627\u0644\u0631\u0635\u064a\u062f", "Paid-leave deduction"),
    "r_lu": ("\u062e\u0635\u0645 \u0627\u0644\u0625\u062c\u0627\u0632\u0629 \u0628\u062f\u0648\u0646 \u0631\u0635\u064a\u062f", "Unpaid-leave deduction"),
    "r_abs": ("\u062e\u0635\u0645 \u0627\u0644\u063a\u064a\u0627\u0628 \u0628\u062f\u0648\u0646 \u0639\u0630\u0631", "Absence deduction"),
    "r_fing": ("\u062e\u0635\u0645 \u062a\u062c\u0627\u0648\u0632 \u0627\u0644\u0628\u0635\u0645\u0627\u062a", "Fingerprint deduction"),
    "r_other": ("\u062e\u0635\u0645\u064a\u0627\u062a \u0623\u062e\u0631\u0649", "Other deductions"),
    "net": ("\u0627\u0644\u0631\u0627\u062a\u0628 \u0627\u0644\u0646\u0647\u0627\u0626\u064a", "Net Salary"),
    "add_h": ("\u0627\u0644\u0625\u0636\u0627\u0641\u0627\u062a", "Additions"),
    "ded_h": ("\u0627\u0644\u062e\u0635\u0645\u064a\u0627\u062a", "Deductions"),
    "btn_xlsx": ("\U0001f4ca  \u062a\u0635\u062f\u064a\u0631 \u0625\u0644\u0649 Excel", "\U0001f4ca  Export to Excel"),
    "btn_print": ("\U0001f5a8  \u0637\u0628\u0627\u0639\u0629 / PDF", "\U0001f5a8  Print / PDF"),
    "btn_save": ("\U0001f4be  \u062d\u0641\u0638 HTML", "\U0001f4be  Save HTML"),
    "saved": ("\u062a\u0645 \u0627\u0644\u062d\u0641\u0638", "Saved"),
    "saved_m": ("\u062a\u0645 \u062d\u0641\u0638 \u0627\u0644\u062a\u0642\u0631\u064a\u0631 \u0628\u0646\u062c\u0627\u062d.", "Report saved successfully."),
    "err": ("\u062e\u0637\u0623", "Error"),
    "hint": ('\u0644\u0644\u062d\u0641\u0638 \u0643\u0640 PDF: \u0627\u0636\u063a\u0637 Ctrl+P \u062b\u0645 \u0627\u062e\u062a\u0631 "\u062d\u0641\u0638 \u0643\u0640 PDF".',
             'To save as PDF: press Ctrl+P then choose "Save as PDF".'),
    "rep_t": ("\u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0631\u0627\u062a\u0628", "Salary Report"),
    "rep_date": ("\u0627\u0644\u062a\u0627\u0631\u064a\u062e", "Date"),
    "xlsx_ok": ("\u062a\u0645 \u0627\u0644\u062a\u0635\u062f\u064a\u0631", "Exported"),
    "xlsx_m": ("\u062a\u0645 \u062a\u0635\u062f\u064a\u0631 \u0646\u062a\u0627\u0626\u062c \u0627\u0644\u0634\u0647\u0631 \u0643\u0634\u064a\u062a \u0628\u0627\u0633\u0645: ",
                "Month results exported as a new sheet named: "),
    "xlsx_need": ("\u0647\u0630\u0647 \u0627\u0644\u0645\u064a\u0632\u0629 \u062a\u062a\u0637\u0644\u0628 \u062d\u0632\u0645\u0629 openpyxl.\n\u062b\u0628\u0651\u062a\u0647\u0627 \u0639\u0628\u0631: pip install openpyxl",
                  "This feature requires the openpyxl package.\nInstall it with: pip install openpyxl"),
    "teos": ("\u0646\u0647\u0627\u064a\u0629 \u0627\u0644\u062e\u062f\u0645\u0629", "End of Service"),
    "s_eos": ("\u0628\u064a\u0627\u0646\u0627\u062a \u0646\u0647\u0627\u064a\u0629 \u0627\u0644\u062e\u062f\u0645\u0629", "End-of-Service Data"),
    "eos_wage": ("\u0627\u0644\u0623\u062c\u0631 \u0627\u0644\u0634\u0647\u0631\u064a \u0627\u0644\u0623\u062e\u064a\u0631 (\u0623\u0633\u0627\u0633\u064a + \u0627\u0644\u0628\u062f\u0644\u0627\u062a)",
                 "Last monthly wage (basic + allowances)"),
    "eos_wage_note": ("\u064a\u064f\u062d\u0633\u0628 \u062a\u0644\u0642\u0627\u0626\u064a\u0627\u064b \u0645\u0646 \u0627\u0644\u0645\u062f\u062e\u0644\u0627\u062a: \u0627\u0644\u0623\u0633\u0627\u0633\u064a + \u0627\u0644\u062b\u0627\u0628\u062a + \u0628\u062f\u0644 \u0627\u0644\u0633\u0643\u0646",
                      "Auto from inputs: basic + fixed + housing"),
    "eos_use": ("\u21bb \u0627\u0633\u062a\u062e\u062f\u0645 \u0627\u0644\u0623\u062c\u0631 \u0645\u0646 \u062a\u0628\u0648\u064a\u0628 \u0627\u0644\u0631\u0627\u062a\u0628",
                "\u21bb Use wage from Salary tab"),
    "eos_start": ("\u062a\u0627\u0631\u064a\u062e \u0627\u0644\u0627\u0644\u062a\u062d\u0627\u0642 (YYYY-MM-DD)", "Joining date (YYYY-MM-DD)"),
    "eos_end": ("\u062a\u0627\u0631\u064a\u062e \u0627\u0646\u062a\u0647\u0627\u0621 \u0627\u0644\u062e\u062f\u0645\u0629 (YYYY-MM-DD)", "End date (YYYY-MM-DD)"),
    "eos_reason": ("\u0633\u0628\u0628 \u0627\u0646\u062a\u0647\u0627\u0621 \u0627\u0644\u062e\u062f\u0645\u0629", "Reason for leaving"),
    "eos_r1": ("\u0627\u0646\u062a\u0647\u0627\u0621 \u0627\u0644\u0639\u0642\u062f / \u0641\u0635\u0644 \u0645\u0646 \u0635\u0627\u062d\u0628 \u0627\u0644\u0639\u0645\u0644",
               "Contract end / employer termination"),
    "eos_r2": ("\u0627\u0633\u062a\u0642\u0627\u0644\u0629", "Resignation"),
    "eos_r3": ("\u062d\u0627\u0644\u0629 \u062e\u0627\u0635\u0629 (\u0627\u0644\u0645\u0627\u062f\u0629 87 \u2014 \u0645\u0643\u0627\u0641\u0623\u0629 \u0643\u0627\u0645\u0644\u0629)",
               "Special case (Art. 87 \u2014 full award)"),
    "eos_res_t": ("\u062a\u0641\u0627\u0635\u064a\u0644 \u0645\u0643\u0627\u0641\u0623\u0629 \u0646\u0647\u0627\u064a\u0629 \u0627\u0644\u062e\u062f\u0645\u0629", "End-of-Service Breakdown"),
    "eos_dur": ("\u0645\u062f\u0629 \u0627\u0644\u062e\u062f\u0645\u0629", "Service duration"),
    "eos_yrs": ("\u0625\u062c\u0645\u0627\u0644\u064a \u0633\u0646\u0648\u0627\u062a \u0627\u0644\u062e\u062f\u0645\u0629", "Total service (years)"),
    "eos_first5": ("\u0623\u0648\u0644 5 \u0633\u0646\u0648\u0627\u062a (\u00bd \u0634\u0647\u0631/\u0633\u0646\u0629)", "First 5 years (\u00bd month/yr)"),
    "eos_after5": ("\u0645\u0627 \u0628\u0639\u062f 5 \u0633\u0646\u0648\u0627\u062a (\u0634\u0647\u0631/\u0633\u0646\u0629)", "After 5 years (1 month/yr)"),
    "eos_full": ("\u0627\u0644\u0645\u0643\u0627\u0641\u0623\u0629 \u0627\u0644\u0643\u0627\u0645\u0644\u0629", "Full award"),
    "eos_factor": ("\u0646\u0633\u0628\u0629 \u0627\u0644\u0627\u0633\u062a\u062d\u0642\u0627\u0642", "Entitlement factor"),
    "eos_total": ("\u0627\u0644\u0645\u0643\u0627\u0641\u0623\u0629 \u0627\u0644\u0645\u0633\u062a\u062d\u0642\u0629", "Entitled Award"),
    "eos_baddate": ("\u062a\u0627\u0631\u064a\u062e \u063a\u064a\u0631 \u0635\u062d\u064a\u062d. \u0627\u0633\u062a\u062e\u062f\u0645 \u0627\u0644\u0635\u064a\u063a\u0629 YYYY-MM-DD.",
                    "Invalid date. Use the format YYYY-MM-DD."),
    "yr": ("\u0633\u0646\u0629", "y"),
    "mo": ("\u0634\u0647\u0631", "m"),
    "dy": ("\u064a\u0648\u0645", "d"),
}


def to_float(s):
    s = "".join(A2W.get(ch, ch) for ch in str(s).strip())
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0


# <<P2>>

NAVY = "#17315c"
ACCENT = "#2f6feb"
ACC_D = "#2457c5"
GREEN = "#15803d"
GREEN_D = "#116a30"
TEAL = "#0f766e"
TEAL_D = "#0b5c55"
RED = "#dc2626"
CARD = "#ffffff"
CANVAS = "#eef2f7"
INK = "#1f2937"
SUB = "#64748b"
LINE = "#e2e8f0"
FIELD = "#f1f5f9"


class SalaryApp(ctk.CTk):
    def __init__(self):
        ctk.set_appearance_mode("light")
        ctk.set_default_color_theme("blue")
        super().__init__()
        self.lang = "ar"
        self.m_month = 8
        self.c_road = "N"
        self.c_house = "Y"
        self.c_contract = "family"
        self.v_basic = ctk.StringVar(value="11625")
        self.v_roadV = ctk.StringVar(value="0")
        self.v_ot = ctk.StringVar(value="0")
        self.v_adv = ctk.StringVar(value="0")
        self.v_med = ctk.StringVar(value="0")
        self.v_lp = ctk.StringVar(value="0")
        self.v_lu = ctk.StringVar(value="11")
        self.v_abs = ctk.StringVar(value="0")
        self.v_fing = ctk.StringVar(value="0")
        self.v_other = ctk.StringVar(value="0")
        self.v_eos_wage = ctk.StringVar(value="11625")
        self.v_eos_start = ctk.StringVar(value="2019-01-01")
        self.v_eos_end = ctk.StringVar(
            value=datetime.date.today().isoformat())
        self.c_eos_reason = "term"
        for v in (self.v_basic, self.v_roadV, self.v_ot, self.v_adv,
                  self.v_med, self.v_lp, self.v_lu, self.v_abs,
                  self.v_fing, self.v_other):
            v.trace_add("write", lambda *a: self.calc())
        for v in (self.v_eos_wage, self.v_eos_start, self.v_eos_end):
            v.trace_add("write", lambda *a: self.calc_eos())
        self.data = {}
        self.eos = {}
        self.res_labels = {}
        self.eos_labels = {}
        self.net_lbl = None
        self.eos_total_lbl = None
        self.body = None
        self.configure(fg_color=CANVAS)
        self.geometry("940x770")
        self.minsize(880, 690)
        self._header()
        self._render()

    # ---------- helpers ----------
    def t(self, k):
        return T[k][0 if self.lang == "ar" else 1]

    def f(self, size=13, bold=False):
        fam = "Tahoma" if self.lang == "ar" else "Segoe UI"
        return ctk.CTkFont(family=fam, size=size,
                           weight="bold" if bold else "normal")

    def _digits(self, s):
        return str(s)

    def fmt(self, x):
        return "{:,.2f}".format(x)

    def _pk(self):
        return ("right", "left") if self.lang == "ar" else ("left", "right")

    def _an(self):
        return ("e", "w") if self.lang == "ar" else ("w", "e")

    # ---------- header (built once) ----------
    def _header(self):
        h = ctk.CTkFrame(self, fg_color=NAVY, corner_radius=0, height=84)
        h.pack(fill="x")
        h.pack_propagate(False)
        self.head_in = ctk.CTkFrame(h, fg_color="transparent")
        self.head_in.pack(fill="both", expand=True, padx=26, pady=16)
        self.title_lbl = ctk.CTkLabel(self.head_in, text="",
                                      text_color="white", font=self.f(23, True))
        self.seg = ctk.CTkSegmentedButton(
            self.head_in, values=["\u0627\u0644\u0639\u0631\u0628\u064a\u0629", "English"],
            command=self._on_lang, font=self.f(13, True),
            fg_color="#24436f", selected_color=ACCENT,
            selected_hover_color=ACC_D, unselected_color="#24436f",
            text_color="white", height=34, corner_radius=8)
        self.seg.set("\u0627\u0644\u0639\u0631\u0628\u064a\u0629")

    def _on_lang(self, value):
        self.lang = "ar" if value == "\u0627\u0644\u0639\u0631\u0628\u064a\u0629" else "en"
        self._render()

    # ---------- render (init + language switch) ----------
    def _render(self):
        self.title(self.t("app"))
        self.title_lbl.configure(text="\U0001f4bc  " + self.t("app"),
                                 font=self.f(23, True))
        self.seg.configure(font=self.f(13, True))
        pa, pb = self._pk()
        self.title_lbl.pack_forget()
        self.seg.pack_forget()
        self.title_lbl.pack(side=pa)
        self.seg.pack(side=pb)
        if self.body is not None:
            self.body.destroy()
        self.body = ctk.CTkFrame(self, fg_color="transparent")
        self.body.pack(fill="both", expand=True, padx=18, pady=16)
        self.tabs = ctk.CTkTabview(
            self.body, fg_color=CARD, corner_radius=14,
            segmented_button_fg_color="#dbe3ef",
            segmented_button_selected_color=ACCENT,
            segmented_button_selected_hover_color=ACC_D,
            segmented_button_unselected_color="#dbe3ef",
            text_color=NAVY, text_color_disabled=SUB)
        self.tabs.pack(fill="both", expand=True)
        self.tabs._segmented_button.configure(font=self.f(14, True))
        tin, tout = self.t("tin"), self.t("tout")
        teos = self.t("teos")
        self.tabs.add(teos)
        self.tabs.add(tout)
        self.tabs.add(tin)
        self.res_labels = {}
        self.eos_labels = {}
        self._build_inputs(self.tabs.tab(tin))
        self._build_results(self.tabs.tab(tout))
        self._build_eos(self.tabs.tab(teos))
        self.tabs.set(tin)
        self.calc()
        self.calc_eos()

    # ---------- small builders ----------
    def _section(self, parent, key):
        card = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=14,
                            border_width=1, border_color=LINE)
        card.pack(fill="x", padx=4, pady=9)
        hd = ctk.CTkFrame(card, fg_color="transparent")
        hd.pack(fill="x", padx=18, pady=(14, 4))
        a, _b = self._pk()
        bar = ctk.CTkFrame(hd, fg_color=ACCENT, width=4, height=20,
                           corner_radius=2)
        bar.pack(side=a, padx=((0, 10) if self.lang != "ar" else (10, 0)))
        bar.pack_propagate(False)
        ctk.CTkLabel(hd, text=self.t(key), text_color=NAVY,
                     font=self.f(15, True)).pack(side=a)
        body = ctk.CTkFrame(card, fg_color="transparent")
        body.pack(fill="x", padx=18, pady=(2, 14))
        return body

    def _row(self, parent, key, make):
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", pady=6)
        an_l, _an_w = self._an()
        lbl = ctk.CTkLabel(row, text=self.t(key), text_color=INK,
                           font=self.f(13), anchor=an_l,
                           justify="right" if self.lang == "ar" else "left",
                           wraplength=360)
        widget = make(row)
        if self.lang == "ar":
            row.grid_columnconfigure(1, weight=1)
            lbl.grid(row=0, column=1, sticky="e")
            widget.grid(row=0, column=0, sticky="w")
        else:
            row.grid_columnconfigure(0, weight=1)
            lbl.grid(row=0, column=0, sticky="w")
            widget.grid(row=0, column=1, sticky="e")

    def _entry(self, parent, var):
        return ctk.CTkEntry(parent, textvariable=var, width=160, height=38,
                            justify="center", font=self.f(14),
                            fg_color=FIELD, border_color=LINE,
                            border_width=1, corner_radius=9)

    def _option(self, parent, values, current, cmd):
        om = ctk.CTkOptionMenu(parent, values=values, command=cmd,
                               width=160, height=38, font=self.f(13),
                               fg_color=ACCENT, button_color=ACC_D,
                               button_hover_color=NAVY, corner_radius=9,
                               dropdown_font=self.f(13))
        om.set(values[current])
        return om

    def _build_inputs(self, tab):
        scroll = ctk.CTkScrollableFrame(tab, fg_color="transparent")
        scroll.pack(fill="both", expand=True)
        scroll.grid_columnconfigure((0, 1), weight=1, uniform="c")
        left = ctk.CTkFrame(scroll, fg_color="transparent")
        right = ctk.CTkFrame(scroll, fg_color="transparent")
        lc, rc = (1, 0) if self.lang == "ar" else (0, 1)
        left.grid(row=0, column=lc, sticky="new", padx=8)
        right.grid(row=0, column=rc, sticky="new", padx=8)

        p = self._section(left, "s_period")
        months = [self._digits(str(i)) for i in range(1, 13)]
        self._row(p, "month", lambda r: self._option(
            r, months, self.m_month - 1,
            lambda v: (setattr(self, "m_month", int(v)), self.calc())))
        self._row(p, "basic", lambda r: self._entry(r, self.v_basic))

        a = self._section(left, "s_allow")
        self._row(a, "road_q", lambda r: self._option(
            r, [self.t("yes"), self.t("no")], 0 if self.c_road == "Y" else 1,
            lambda v: (setattr(self, "c_road",
                               "Y" if v == self.t("yes") else "N"),
                       self.calc())))
        self._row(a, "road_v", lambda r: self._entry(r, self.v_roadV))
        self._row(a, "house_q", lambda r: self._option(
            r, [self.t("yes"), self.t("no")], 0 if self.c_house == "Y" else 1,
            lambda v: (setattr(self, "c_house",
                               "Y" if v == self.t("yes") else "N"),
                       self.calc())))
        self._row(a, "contract", lambda r: self._option(
            r, [self.t("c_fam"), self.t("c_sin")],
            0 if self.c_contract == "family" else 1,
            lambda v: (setattr(self, "c_contract",
                               "family" if v == self.t("c_fam")
                               else "single"), self.calc())))
        self._row(a, "ot", lambda r: self._entry(r, self.v_ot))

        d = self._section(right, "s_ded")
        self._row(d, "adv", lambda r: self._entry(r, self.v_adv))
        self._row(d, "med", lambda r: self._entry(r, self.v_med))

        g = self._section(right, "s_leave")
        self._row(g, "lp", lambda r: self._entry(r, self.v_lp))
        self._row(g, "lu", lambda r: self._entry(r, self.v_lu))
        self._row(g, "absn", lambda r: self._entry(r, self.v_abs))
        self._row(g, "fing", lambda r: self._entry(r, self.v_fing))
        self._row(g, "other", lambda r: self._entry(r, self.v_other))


# <<P3>>

    # ---------- results ----------
    def _res_row(self, parent, key, color):
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", pady=3)
        an_l, _w = self._an()
        lbl = ctk.CTkLabel(row, text=self.t(key), text_color=INK,
                           font=self.f(13), anchor=an_l)
        val = ctk.CTkLabel(row, text="-", text_color=color,
                           font=self.f(14, True))
        if self.lang == "ar":
            row.grid_columnconfigure(1, weight=1)
            lbl.grid(row=0, column=1, sticky="e")
            val.grid(row=0, column=0, sticky="w")
        else:
            row.grid_columnconfigure(0, weight=1)
            lbl.grid(row=0, column=0, sticky="w")
            val.grid(row=0, column=1, sticky="e")
        self.res_labels[key] = val

    def _col_head(self, parent, key, color):
        an_l, _w = self._an()
        ctk.CTkLabel(parent, text=self.t(key), text_color=color,
                     font=self.f(14, True), anchor=an_l).pack(
                         fill="x", pady=(0, 4))

    def _build_results(self, tab):
        bottom = ctk.CTkFrame(tab, fg_color="transparent")
        bottom.pack(side="bottom", fill="x", padx=6, pady=(8, 6))
        self.net_lbl = ctk.CTkLabel(bottom, text="", fg_color=NAVY,
                                    text_color="white", corner_radius=12,
                                    height=60, font=self.f(21, True))
        self.net_lbl.pack(fill="x", pady=(0, 10))
        ctk.CTkButton(bottom, text=self.t("btn_xlsx"),
                      command=self.export_excel, fg_color=TEAL,
                      hover_color=TEAL_D, height=48, corner_radius=10,
                      font=self.f(16, True)).pack(fill="x", pady=(0, 8))
        bb = ctk.CTkFrame(bottom, fg_color="transparent")
        bb.pack(fill="x")
        bb.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(bb, text=self.t("btn_print"),
                      command=lambda: self.save_report(True),
                      fg_color=GREEN, hover_color=GREEN_D, height=42,
                      corner_radius=10, font=self.f(13, True)).grid(
                          row=0, column=0, sticky="ew", padx=(0, 5))
        ctk.CTkButton(bb, text=self.t("btn_save"),
                      command=lambda: self.save_report(False),
                      fg_color=ACCENT, hover_color=ACC_D, height=42,
                      corner_radius=10, font=self.f(13, True)).grid(
                          row=0, column=1, sticky="ew", padx=(5, 0))

        scroll = ctk.CTkScrollableFrame(tab, fg_color="transparent")
        scroll.pack(fill="both", expand=True)

        card = ctk.CTkFrame(scroll, fg_color=CARD, corner_radius=14,
                            border_width=1, border_color=LINE)
        card.pack(fill="x", padx=4, pady=4)
        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=22, pady=18)

        hd = ctk.CTkFrame(inner, fg_color="transparent")
        hd.pack(fill="x")
        a, _b = self._pk()
        bar = ctk.CTkFrame(hd, fg_color=ACCENT, width=5, height=24,
                           corner_radius=2)
        bar.pack(side=a, padx=((0, 10) if self.lang != "ar" else (10, 0)))
        bar.pack_propagate(False)
        ctk.CTkLabel(hd, text=self.t("res_t"), text_color=NAVY,
                     font=self.f(17, True)).pack(side=a)

        info = ctk.CTkFrame(inner, fg_color="transparent")
        info.pack(fill="x", pady=(12, 2))
        for k in ("r_days", "r_work", "r_day", "r_hour"):
            self._res_row(info, k, SUB)

        ctk.CTkFrame(inner, fg_color=LINE, height=1).pack(fill="x", pady=12)

        cols = ctk.CTkFrame(inner, fg_color="transparent")
        cols.pack(fill="both", expand=True)
        cols.grid_columnconfigure((0, 1), weight=1, uniform="r")
        addc = ctk.CTkFrame(cols, fg_color="transparent")
        dedc = ctk.CTkFrame(cols, fg_color="transparent")
        ac, dc = (1, 0) if self.lang == "ar" else (0, 1)
        addc.grid(row=0, column=ac, sticky="new", padx=12)
        dedc.grid(row=0, column=dc, sticky="new", padx=12)

        self._col_head(addc, "add_h", GREEN)
        for k in ("r_basic", "r_fixed", "r_road", "r_house", "r_ot"):
            self._res_row(addc, k, GREEN)
        self._col_head(dedc, "ded_h", RED)
        for k in ("r_adv", "r_med", "r_lp", "r_lu", "r_abs", "r_fing",
                  "r_other"):
            self._res_row(dedc, k, RED)

    # ---------- logic (same formulas as the Excel) ----------
    def calc(self, *args):
        if not self.res_labels:
            return
        m = min(max(self.m_month, 1), 12)
        A5 = DAYS[m - 1]
        A6 = FRI[m - 1]
        A7 = A5 - A6
        A9 = to_float(self.v_basic.get())
        A10 = A9 * 0.333333333333333
        A11 = A9 / A5 if A5 else 0
        A12 = A11 / 8
        A14 = to_float(self.v_roadV.get()) if self.c_road == "Y" else 0.0
        if A14 in (0.07, 0.15):
            A15 = (A9 + A10) * A14
        else:
            A15 = A14 * 1.5 * A7 * A12
        A19 = 3 if self.c_contract == "family" else 2
        A20 = 0.0 if self.c_house != "Y" else (A9 * A19 / 12)
        A22 = to_float(self.v_ot.get())
        A23 = A12 * 1.5 * A22
        A25 = to_float(self.v_adv.get())
        A26 = 0.0 if A25 == 0 else 0.15 * (A9 + A10)
        A27 = to_float(self.v_med.get())
        A28 = 0.0 if A27 == 0 else 0.15 * (A9 + A10)
        A30 = to_float(self.v_lp.get())
        A31 = (A15 / A5) * A30 if A5 else 0
        A32 = to_float(self.v_lu.get())
        A33 = A32 * ((A9 + A10 + A20) / A5) if A5 else 0
        A34 = to_float(self.v_abs.get())
        A35 = ((A34 / A5) * (A9 + A10 + A15)) * 2 if A5 else 0
        A36 = to_float(self.v_fing.get())
        A37 = ((((A9 + A10 + A15 + A20) / A5) / 10) * 2.5 * A36) if A5 else 0
        A38 = to_float(self.v_other.get())
        A39 = (A9 + A10 + A15 + A20 + A23
               - A26 - A28 - A31 - A33 - A35 - A37 - A38)

        self.data = {
            "month": m, "days": A5, "work": A7, "day": A11, "hour": A12,
            "r_basic": A9, "r_fixed": A10, "r_road": A15, "r_house": A20,
            "r_ot": A23, "r_adv": A26, "r_med": A28, "r_lp": A31,
            "r_lu": A33, "r_abs": A35, "r_fing": A37, "r_other": A38,
            "net": A39,
        }
        R = self.res_labels
        R["r_days"].configure(text=self._digits(A5))
        R["r_work"].configure(text=self._digits(A7))
        R["r_day"].configure(text=self.fmt(A11))
        R["r_hour"].configure(text=self.fmt(A12))
        for k in ("r_basic", "r_fixed", "r_road", "r_house", "r_ot",
                  "r_adv", "r_med", "r_lp", "r_lu", "r_abs", "r_fing",
                  "r_other"):
            R[k].configure(text=self.fmt(self.data[k]))
        if self.net_lbl is not None:
            self.net_lbl.configure(
                text=self.t("net") + ":   " + self.fmt(A39))

        self.v_eos_wage.set(str(round(A9 + A10 + A20, 2)))

    # ---------- End-of-Service (Saudi Labor Law) ----------
    def _eos_row(self, parent, key, color, big=False):
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", pady=3)
        an_l, _w = self._an()
        lbl = ctk.CTkLabel(row, text=self.t(key), text_color=INK,
                           font=self.f(14 if big else 13, big), anchor=an_l)
        val = ctk.CTkLabel(row, text="-", text_color=color,
                           font=self.f(15 if big else 14, True))
        if self.lang == "ar":
            row.grid_columnconfigure(1, weight=1)
            lbl.grid(row=0, column=1, sticky="e")
            val.grid(row=0, column=0, sticky="w")
        else:
            row.grid_columnconfigure(0, weight=1)
            lbl.grid(row=0, column=0, sticky="w")
            val.grid(row=0, column=1, sticky="e")
        self.eos_labels[key] = val

    def _eos_reason_set(self, v):
        if v == self.t("eos_r2"):
            self.c_eos_reason = "resign"
        elif v == self.t("eos_r3"):
            self.c_eos_reason = "special"
        else:
            self.c_eos_reason = "term"
        self.calc_eos()

    def _eos_use_wage(self):
        d = self.data or {}
        total = (d.get("r_basic", 0) + d.get("r_fixed", 0)
                 + d.get("r_house", 0))
        self.v_eos_wage.set(str(round(total, 2)))

    def _entry_ro(self, parent, var):
        e = ctk.CTkEntry(parent, textvariable=var, width=160, height=38,
                         justify="center", font=self.f(14),
                         fg_color=FIELD, border_color=LINE,
                         border_width=1, corner_radius=9,
                         text_color=NAVY)
        e.configure(state="disabled")
        return e

    def _build_eos(self, tab):
        scroll = ctk.CTkScrollableFrame(tab, fg_color="transparent")
        scroll.pack(fill="both", expand=True)

        s = self._section(scroll, "s_eos")
        self._row(s, "eos_wage",
                  lambda r: self._entry_ro(r, self.v_eos_wage))
        ctk.CTkLabel(s, text=self.t("eos_wage_note"), text_color=SUB,
                     font=self.f(11), anchor=self._an()[0],
                     justify="right" if self.lang == "ar" else "left",
                     wraplength=380).pack(
                         fill="x", pady=(0, 8),
                         padx=(0, 4) if self.lang != "ar" else (4, 0))
        self._row(s, "eos_start",
                  lambda r: self._entry(r, self.v_eos_start))
        self._row(s, "eos_end", lambda r: self._entry(r, self.v_eos_end))
        self._row(s, "eos_reason", lambda r: self._option(
            r, [self.t("eos_r1"), self.t("eos_r2"), self.t("eos_r3")],
            {"term": 0, "resign": 1, "special": 2}[self.c_eos_reason],
            self._eos_reason_set))

        card = ctk.CTkFrame(scroll, fg_color=CARD, corner_radius=14,
                            border_width=1, border_color=LINE)
        card.pack(fill="x", padx=4, pady=9)
        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=22, pady=18)
        hd = ctk.CTkFrame(inner, fg_color="transparent")
        hd.pack(fill="x")
        a, _b = self._pk()
        bar = ctk.CTkFrame(hd, fg_color=ACCENT, width=5, height=24,
                           corner_radius=2)
        bar.pack(side=a, padx=((0, 10) if self.lang != "ar" else (10, 0)))
        bar.pack_propagate(False)
        ctk.CTkLabel(hd, text=self.t("eos_res_t"), text_color=NAVY,
                     font=self.f(17, True)).pack(side=a)

        info = ctk.CTkFrame(inner, fg_color="transparent")
        info.pack(fill="x", pady=(12, 2))
        self._eos_row(info, "eos_dur", SUB)
        self._eos_row(info, "eos_yrs", SUB)

        ctk.CTkFrame(inner, fg_color=LINE, height=1).pack(fill="x", pady=12)

        det = ctk.CTkFrame(inner, fg_color="transparent")
        det.pack(fill="x", pady=(0, 2))
        self._eos_row(det, "eos_first5", GREEN)
        self._eos_row(det, "eos_after5", GREEN)
        self._eos_row(det, "eos_full", INK)
        self._eos_row(det, "eos_factor", ACCENT)

        self.eos_total_lbl = ctk.CTkLabel(
            scroll, text="", fg_color=NAVY, text_color="white",
            corner_radius=12, height=60, font=self.f(21, True))
        self.eos_total_lbl.pack(fill="x", padx=4, pady=(2, 12))

    def _parse_date(self, s):
        s = "".join(A2W.get(ch, ch) for ch in str(s).strip())
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    def _date_diff(self, d1, d2):
        import calendar
        y = d2.year - d1.year
        m = d2.month - d1.month
        dd = d2.day - d1.day
        if dd < 0:
            m -= 1
            pm = d2.month - 1 or 12
            py = d2.year if d2.month > 1 else d2.year - 1
            dd += calendar.monthrange(py, pm)[1]
        if m < 0:
            y -= 1
            m += 12
        return y, m, dd

    def calc_eos(self, *args):
        if not self.eos_labels:
            return
        wage = to_float(self.v_eos_wage.get())
        d1 = self._parse_date(self.v_eos_start.get())
        d2 = self._parse_date(self.v_eos_end.get())
        if d1 is None or d2 is None or d2 < d1:
            for lbl in self.eos_labels.values():
                lbl.configure(text="\u2014")
            if self.eos_total_lbl is not None:
                self.eos_total_lbl.configure(
                    text=self.t("eos_baddate"), font=self.f(14, True))
            self.eos = {}
            return
        total_days = (d2 - d1).days
        y, m, dd = self._date_diff(d1, d2)
        Y = total_days / 365.0
        first5 = wage * 0.5 * min(Y, 5.0)
        after5 = wage * 1.0 * max(0.0, Y - 5.0)
        full = first5 + after5
        if self.c_eos_reason == "resign":
            if Y < 2:
                factor = 0.0
            elif Y < 5:
                factor = 1.0 / 3.0
            elif Y < 10:
                factor = 2.0 / 3.0
            else:
                factor = 1.0
        else:
            factor = 1.0
        total = full * factor
        self.eos = {
            "wage": wage, "days": total_days, "years": Y,
            "first5": first5, "after5": after5, "full": full,
            "factor": factor, "total": total,
            "ymd": (y, m, dd),
        }
        dur = (str(y) + " " + self.t("yr") + "  " + str(m) + " "
               + self.t("mo") + "  " + str(dd) + " " + self.t("dy"))
        self.eos_labels["eos_dur"].configure(text=dur)
        self.eos_labels["eos_yrs"].configure(text="{:,.2f}".format(Y))
        self.eos_labels["eos_first5"].configure(text=self.fmt(first5))
        self.eos_labels["eos_after5"].configure(text=self.fmt(after5))
        self.eos_labels["eos_full"].configure(text=self.fmt(full))
        self.eos_labels["eos_factor"].configure(
            text="{:.1f}%".format(factor * 100))
        if self.eos_total_lbl is not None:
            self.eos_total_lbl.configure(
                text=self.t("eos_total") + ":   " + self.fmt(total),
                font=self.f(21, True))

# <<P4>>

    # ---------- report (print / PDF via browser) ----------
    def _report_html(self):
        d = self.data
        now = datetime.datetime.now().strftime("%Y-%m-%d  %H:%M")
        rtl = "rtl" if self.lang == "ar" else "ltr"
        ff = "Tahoma" if self.lang == "ar" else "'Segoe UI'"
        add = ["r_basic", "r_fixed", "r_road", "r_house", "r_ot"]
        ded = ["r_adv", "r_med", "r_lp", "r_lu", "r_abs", "r_fing", "r_other"]
        rows = ""
        for k in add:
            rows += ("<tr><td>" + html.escape(self.t(k))
                     + "</td><td class='p'>+ " + self.fmt(d[k])
                     + "</td></tr>")
        for k in ded:
            rows += ("<tr><td>" + html.escape(self.t(k))
                     + "</td><td class='m'>- " + self.fmt(d[k])
                     + "</td></tr>")
        head = (
            "<!DOCTYPE html><html lang='" + self.lang + "' dir='" + rtl
            + "'><head><meta charset='utf-8'><title>" + self.t("rep_t")
            + "</title><style>*{font-family:" + ff + ",Arial,sans-serif}"
            "body{padding:28px;color:#222;max-width:720px;margin:auto}"
            "h1{color:#17315c;text-align:center;margin:0 0 4px}"
            ".meta{text-align:center;color:#777;font-size:13px;"
            "margin-bottom:16px}"
            "table{width:100%;border-collapse:collapse;font-size:14px}"
            "td{padding:9px 12px;border-bottom:1px solid #e5e8f0}"
            "td:last-child{text-align:" + ("left" if self.lang == "ar"
                                           else "right")
            + ";font-weight:700;white-space:nowrap}"
            ".p{color:#15803d}.m{color:#dc2626}"
            ".net{margin-top:16px;background:#17315c;color:#fff;padding:16px;"
            "border-radius:12px;display:flex;justify-content:space-between;"
            "font-size:20px;font-weight:700}"
            "@media print{.hint{display:none}}"
            ".hint{margin-top:14px;color:#888;font-size:12px;"
            "text-align:center}</style></head><body>")
        meta = ("<h1>" + self.t("rep_t") + "</h1><div class='meta'>"
                + self.t("month") + ": " + self._digits(d["month"])
                + " &nbsp;|&nbsp; " + self.t("rep_date") + ": "
                + html.escape(now) + "</div>")
        net = ("<div class='net'><span>" + self.t("net") + "</span><span>"
               + self.fmt(d["net"]) + "</span></div>")
        hint = "<div class='hint'>" + html.escape(self.t("hint")) \
               + "</div></body></html>"
        return head + meta + "<table>" + rows + "</table>" + net + hint

    def save_report(self, open_now=True):
        try:
            doc = self._report_html()
            if open_now:
                path = os.path.join(tempfile.gettempdir(),
                                    "salary_report.html")
            else:
                path = filedialog.asksaveasfilename(
                    defaultextension=".html",
                    filetypes=[("HTML", "*.html")],
                    initialfile="salary_report.html")
                if not path:
                    return
            with open(path, "w", encoding="utf-8") as fp:
                fp.write(doc)
            if open_now:
                webbrowser.open("file://" + os.path.abspath(path))
            else:
                messagebox.showinfo(self.t("saved"), self.t("saved_m"))
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror(self.t("err"), str(exc))

    # ---------- export to Excel (one sheet per export, named by date) ----
    def export_excel(self):
        if not _HAS_XLSX:
            messagebox.showwarning("Excel", self.t("xlsx_need"))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="salary_history.xlsx",
            confirmoverwrite=False)
        if not path:
            return
        try:
            if os.path.exists(path):
                wb = load_workbook(path)
            else:
                wb = Workbook()
                wb.remove(wb.active)
            base = datetime.date.today().isoformat()
            name, i = base, 2
            while name in wb.sheetnames:
                name = base + " (" + str(i) + ")"
                i += 1
            ws = wb.create_sheet(title=name)
            wb.move_sheet(ws, offset=-(len(wb.sheetnames) - 1))
            wb.active = wb.sheetnames.index(name)
            self._fill_sheet(ws)
            wb.save(path)
            messagebox.showinfo(self.t("xlsx_ok"), self.t("xlsx_m") + name)
        except PermissionError:
            messagebox.showerror(
                self.t("err"),
                "Excel file is open. Close it and try again. / "
                "\u0645\u0644\u0641 \u0627\u0644\u0625\u0643\u0633\u0644 \u0645\u0641\u062a\u0648\u062d\u060c \u0623\u063a\u0644\u0642\u0647 \u062b\u0645 \u0623\u0639\u062f \u0627\u0644\u0645\u062d\u0627\u0648\u0644\u0629.")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror(self.t("err"), str(exc))

    def _fill_sheet(self, ws):
        d = self.data
        navy = "FF17315C"
        green = "FF15803D"
        red = "FFDC2626"
        thin = Side(style="thin", color="FFDFE3EB")
        border = Border(bottom=thin)
        money = "#,##0.00"
        if self.lang == "ar":
            ws.sheet_view.rightToLeft = True
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 20

        def put(r, label, value, color=None, fmt=None, fill=None,
                bold=False, size=11, white=False):
            c1 = ws.cell(row=r, column=1, value=label)
            c1.font = Font(name="Segoe UI", size=size, bold=bold,
                           color="FFFFFFFF" if white else "FF1F2937")
            c1.alignment = Alignment(horizontal="right" if self.lang == "ar"
                                     else "left", vertical="center")
            c2 = ws.cell(row=r, column=2, value=value)
            c2.font = Font(name="Segoe UI", size=size, bold=bold,
                           color=color or ("FFFFFFFF" if white
                                           else "FF1F2937"))
            c2.alignment = Alignment(horizontal="center", vertical="center")
            if fmt:
                c2.number_format = fmt
            if fill:
                pf = PatternFill("solid", fgColor=fill)
                c1.fill = pf
                c2.fill = pf
            c1.border = border
            c2.border = border
            ws.row_dimensions[r].height = 20

        def header(r, text, fill):
            ws.merge_cells(start_row=r, start_column=1,
                           end_row=r, end_column=2)
            c = ws.cell(row=r, column=1, value=text)
            c.font = Font(name="Segoe UI", size=12, bold=True,
                          color="FFFFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.fill = PatternFill("solid", fgColor=fill)
            ws.row_dimensions[r].height = 24

        ws.merge_cells("A1:B1")
        tc = ws.cell(row=1, column=1, value=self.t("rep_t"))
        tc.font = Font(name="Segoe UI", size=15, bold=True,
                       color="FFFFFFFF")
        tc.alignment = Alignment(horizontal="center", vertical="center")
        tc.fill = PatternFill("solid", fgColor=navy)
        ws.row_dimensions[1].height = 30

        put(3, self.t("month"), d["month"])
        put(4, self.t("rep_date"),
            datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
        put(5, self.t("r_days"), d["days"])
        put(6, self.t("r_work"), d["work"])
        put(7, self.t("r_day"), round(d["day"], 2), fmt=money)
        put(8, self.t("r_hour"), round(d["hour"], 2), fmt=money)

        header(10, self.t("add_h"), green)
        r = 11
        for k in ("r_basic", "r_fixed", "r_road", "r_house", "r_ot"):
            put(r, self.t(k), round(d[k], 2), color=green, fmt=money)
            r += 1
        r += 1
        header(r, self.t("ded_h"), red)
        r += 1
        for k in ("r_adv", "r_med", "r_lp", "r_lu", "r_abs", "r_fing",
                  "r_other"):
            put(r, self.t(k), round(d[k], 2), color=red, fmt=money)
            r += 1
        r += 1
        put(r, self.t("net"), round(d["net"], 2), fmt=money,
            fill=navy, bold=True, size=13, white=True)
        ws.row_dimensions[r].height = 26


def _log_path():
    try:
        if getattr(sys, "frozen", False):
            base = os.path.dirname(sys.executable)
        else:
            base = os.path.dirname(os.path.abspath(__file__))
        return os.path.join(base, "salary_error.log")
    except Exception:
        return os.path.join(tempfile.gettempdir(), "salary_error.log")


if __name__ == "__main__":
    try:
        SalaryApp().mainloop()
    except Exception:
        err = traceback.format_exc()
        try:
            with open(_log_path(), "w", encoding="utf-8") as fp:
                fp.write(err)
        except Exception:
            pass
        try:
            messagebox.showerror(
                "Error \u2014 \u062e\u0637\u0623",
                "\u062d\u062f\u062b \u062e\u0637\u0623 \u0623\u062b\u0646\u0627\u0621 \u0627\u0644\u062a\u0634\u063a\u064a\u0644:\n\n"
                + err[-1600:])
        except Exception:
            pass
        sys.exit(1)
